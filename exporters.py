"""
exporters.py
------------
Turns the transcript + summary text into downloadable Word (.docx) and
PDF files, and sends them by email.
"""

import html
import os
import smtplib
from email.message import EmailMessage
from pathlib import Path


# --------------------------------------------------------------------------- #
# Word (.docx) export
# --------------------------------------------------------------------------- #
def build_docx(summary: str, transcript: str, lang_code: str, lang_prob: str, out_path: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    title = doc.add_heading("Voice File Summary", level=1)

    meta = doc.add_paragraph()
    meta.add_run(f"Detected source language: {lang_code} (confidence: {lang_prob})").italic = True

    doc.add_heading("Summary (English)", level=2)
    for para in summary.split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_heading("Full Transcript (English)", level=2)
    for para in transcript.split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.save(out_path)


# --------------------------------------------------------------------------- #
# PDF export
# --------------------------------------------------------------------------- #
def build_pdf(summary: str, transcript: str, lang_code: str, lang_prob: str, out_path: str) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import inch
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    doc = SimpleDocTemplate(
        out_path,
        pagesize=LETTER,
        leftMargin=0.85 * inch,
        rightMargin=0.85 * inch,
        topMargin=0.85 * inch,
        bottomMargin=0.85 * inch,
        title="Voice File Summary",
    )

    styles = getSampleStyleSheet()
    meta_style = ParagraphStyle(
        "Meta", parent=styles["Normal"], fontSize=9, textColor="#555555", spaceAfter=14,
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontSize=11, leading=16, spaceAfter=8,
    )

    story = [
        Paragraph("Voice File Summary", styles["Title"]),
        Paragraph(
            f"Detected source language: {_pdf_safe(lang_code)} (confidence: {_pdf_safe(lang_prob)})",
            meta_style,
        ),
        Paragraph("Summary (English)", styles["Heading2"]),
    ]
    for para in summary.split("\n"):
        if para.strip():
            story.append(Paragraph(_pdf_safe(para.strip()), body_style))
    story.append(Spacer(1, 10))

    story.append(Paragraph("Full Transcript (English)", styles["Heading2"]))
    for para in transcript.split("\n"):
        if para.strip():
            story.append(Paragraph(_pdf_safe(para.strip()), body_style))

    doc.build(story)


def _pdf_safe(text: str) -> str:
    """
    Prepares text for the PDF: the built-in fonts only support latin-1,
    so anything outside that range is replaced with '?'; the result is
    then HTML-escaped since ReportLab's Paragraph accepts a small subset
    of HTML markup and would otherwise choke on stray & < > characters.
    """
    latin1_safe = text.encode("latin-1", errors="replace").decode("latin-1")
    return html.escape(latin1_safe)


# --------------------------------------------------------------------------- #
# Email delivery
# --------------------------------------------------------------------------- #
def send_summary_email(
    to_address: str,
    subject: str,
    body: str,
    attachment_paths: list[str],
) -> None:
    """
    Sends the summary (and optional attachments) by email using SMTP
    credentials from environment variables:

        SMTP_HOST      e.g. smtp.gmail.com
        SMTP_PORT      e.g. 587
        SMTP_USER      your email address / login
        SMTP_PASSWORD  your email password or app password
        SMTP_FROM      (optional) "From" address, defaults to SMTP_USER

    For Gmail specifically, you must use an "App Password", not your
    regular Google password (Google Account -> Security -> 2-Step
    Verification -> App passwords).

    If you'd rather not deal with environment variables, you can hardcode
    your credentials below instead — just replace the os.environ.get(...)
    calls for user/password with plain strings, e.g.:
        user = "you@gmail.com"
        password = "your-app-password"
    Only do this on your own machine, and never commit real credentials
    to a shared or public repository.
    """
    host = os.environ.get("SMTP_HOST") or None
    # "or" (rather than a get() default) so an empty value in .env, like a
    # bare "SMTP_PORT=" or "SMTP_FROM=" line, counts as unset instead of
    # being used as-is.
    port = int(os.environ.get("SMTP_PORT") or "587")
    user = os.environ.get("SMTP_USER") or None
    password = os.environ.get("SMTP_PASSWORD") or None
    from_addr = os.environ.get("SMTP_FROM") or user

    if not all([host, user, password]):
        raise RuntimeError(
            "Email is not configured on the server. Set SMTP_HOST, SMTP_USER, "
            "and SMTP_PASSWORD environment variables (see README)."
        )

    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = from_addr
    msg["To"] = to_address
    msg.set_content(body)

    for path in attachment_paths:
        p = Path(path)
        if not p.exists():
            continue
        data = p.read_bytes()
        if p.suffix == ".docx":
            maintype, subtype = (
                "application",
                "vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        elif p.suffix == ".pdf":
            maintype, subtype = "application", "pdf"
        else:
            maintype, subtype = "text", "plain"
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=p.name)

    # Port 465 uses implicit SSL (encrypted from the first byte); port 587
    # starts plain and upgrades with STARTTLS. Some networks block one but
    # not the other, so both are supported. A timeout keeps a blocked
    # connection from hanging the request for minutes.
    if port == 465:
        with smtplib.SMTP_SSL(host, port, timeout=20) as server:
            server.login(user, password)
            server.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=20) as server:
            server.starttls()
            server.login(user, password)
            server.send_message(msg)
